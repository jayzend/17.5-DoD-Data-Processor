import pandas as pd
import numpy as np
import openpyxl
import streamlit as st
import io

st.set_page_config(page_title="Battery Data Processor", layout="wide")
st.title("Universal Battery Data Processor")
st.write("Upload your raw test data to generate the processed Excel file and charts for DCA and DoD tests.")

# --- Step 1: Input Setup ---
uploaded_file = st.file_uploader("Select Raw Data Excel File", type=["xlsx"])

if uploaded_file is not None:
    with st.spinner("Analyzing and Processing Data..."):
        try:
            # --- Step 2: Test Type & Sheet Detection ---
            xl = pd.ExcelFile(uploaded_file, engine='openpyxl')
            all_sheets = xl.sheet_names

            # --- Step 3: Filtering of Sheet ---
            target_sheet = None
            test_type = None
            is_failed = False

            for s in all_sheets:
                s_lower = s.lower()
                
                if "failed" in s_lower:
                    is_failed = True
                    
                if "17.5" in s_lower:
                    test_type = "DoD_17.5"
                    target_sheet = s
                    break

                elif "50" in s_lower:
                    test_type = "DoD_50"
                    target_sheet = s
                    break

                elif "dca" in s_lower and 'flooded' not in s_lower:
                    test_type = "DCA"
                    target_sheet = s
                    break
                    
                elif "mht" in s_lower:
                    test_type = "MHT"
                    target_sheet = s
                    break

            if not target_sheet:
                st.error("No valid test sheet was identified in the workbook.")
                st.stop()

            # --- Step 4: Read Raw Data ---
            df_meta = pd.read_excel(xl, sheet_name=target_sheet, header=None, nrows=4)
            if df_meta.shape[0] >= 4 and df_meta.shape[1] >= 2:
                battery_code = str(df_meta.iloc[3, 1]).strip()
            else:
                battery_code = "UNKNOWN"

            raw_df = pd.read_excel(xl, sheet_name=target_sheet, skiprows=29)
            raw_df.columns = raw_df.columns.astype(str).str.strip()

            status_text = " - Failed" if is_failed else ""
            output_filename = f'Processed {test_type}{status_text} Battery Results of {battery_code}.xlsx'

            # --- 3. Create the Output File in Memory ---
            output_buffer = io.BytesIO()

            # --- Step 5: Execute Test ---
            with pd.ExcelWriter(output_buffer, engine='xlsxwriter') as writer:
                workbook = writer.book
                
                # Test A: DCA TEST ENGINE
                if test_type == "DCA":
                    # --- IC Sheet Filtering and Computation ---
                    df = raw_df.copy()
                    df = df[df['Status'] == 'CHA']
                    df = df[df['Step'].isin([30])]
                    df = df[df['Voltage'] >= 14.8]
                    df = df[df['Current'] >= 0]
                    df = df[(df['Step Time'] = '00:00:10.000') & (df['Step Time'] < '00:00:10.1')]
                    df = df.dropna(subset=['Temperature'])
                    df['Ampere'] = df['AhStep'] * 360
                    df['A/Ah'] = df['Ampere'] / 60

                    total_ah1 = df['AhStep'].sum()
                    total_as1 = total_ah1 * 3600
                    time_sec1 = 200
                    amp1 = total_as1 / time_sec1
                    capacity1 = 60
                    a_ah1 = amp1 / capacity1
                    factor1 = 0.512
                    net1 = a_ah1 * factor1

                    df_summary = pd.DataFrame({
                        'Metric': ['Total-Ah', 'Total-AS', 'Time-Sec', 'Amp', 'Capacity-Ah', 'A/Ah', 'Factor', 'Net'],
                        'Value':  [total_ah1, total_as1, time_sec1, amp1, capacity1, a_ah1, factor1, net1]
                    })

                    # --- ID Sheet Filtering and Computation ---
                    df2 = raw_df.copy()
                    df2 = df2[df2['Status'] == 'CHA']
                    df2 = df2[df2['Step'].isin([43])]
                    df2 = df2[df2['Voltage'] >= 14.8]
                    df2 = df2[df2['Current'] >= 0]
                    df2 = df2[(df2['Step Time'] = '00:00:10.000') & (df2['Step Time'] < '00:00:10.1')]
                    df2 = df2.dropna(subset=['Temperature'])
                    df2['Ampere'] = df2['AhStep'] * 360
                    df2['A/Ah'] = df2['Ampere'] / 60

                    total_ah2 = df2['AhStep'].sum()
                    total_as2 = total_ah2 * 3600
                    time_sec2 = 200
                    amp2 = total_as2 / time_sec2
                    capacity2 = 60
                    a_ah2 = amp2 / capacity2
                    factor2 = 0.223
                    net2 = a_ah2 * factor2

                    df2_summary = pd.DataFrame({
                        'Metric': ['Total-Ah', 'Total-AS', 'Time-Sec', 'Amp', 'Capacity-Ah', 'A/Ah', 'Factor', 'Net'],
                        'Value':  [total_ah2, total_as2, time_sec2, amp2, capacity2, a_ah2, factor2, net2]
                    })

                    # --- IR Sheet and Computation ---
                    df3 = raw_df.copy()
                    df3 = df3[df3['Status'] == 'CHA']
                    df3 = df3[df3['Step'].isin([70, 81])]
                    df3 = df3[df3['Cycle Level'] == 3]
                    df3 = df3[(df3['Step Time'] = '00:00:05.000') & (df3['Step Time'] < '00:00:05.1')]
                    df3 = df3.dropna(subset=['Temperature'])
                    
                    df3['Time Stamp'] = df3['Time Stamp'].astype(str).str.strip()
                    df3['Status'] = df3['Status'].astype(str).str.strip()
                    df3['Step'] = df3['Step'].astype(str).str.strip()
                    df3 = df3.drop_duplicates(subset=['Time Stamp', 'Step', 'Status'], keep='first')
                    
                    df3['Ampere'] = df3['AhStep'] * 720
                    df3['A/Ah'] = df3['Ampere'] / 60

                    # --- Tracker --- 
                    tracker = (df3['Cycle'] < df3['Cycle'].shift(1)).cumsum() + 1
                    df3['S'] = 'S' + tracker.astype(str)

                    # --- IR Calculation ---
                    IR_Calculation = df3.groupby('S', sort=False)['AhStep'].sum().reset_index()
                    IR_Calculation = IR_Calculation.rename(columns={'AhStep': 'Total AhStep'})

                    grand_total_ah = df3['AhStep'].sum()
                    total_as3 = grand_total_ah * 3600
                    time_sec3 = 2850
                    amp3 = total_as3 / time_sec3
                    capacity3 = 60
                    a_ah3 = amp3 / capacity3
                    factor3 = 0.218
                    net3 = a_ah3 * factor3

                    df4 = pd.DataFrame({
                        '1': ['Total-Ah', 'Total-AS', 'Time-Sec', 'Amp', 'Capacity-Ah', 'A/Ah', 'Factor', 'Net'],
                        '2': [grand_total_ah, total_as3, time_sec3, amp3, capacity3, a_ah3, factor3, net3]
                    })

                    Ic, Id, Ir, Constant = net1, net2, net3, 0.181
                    final_a_ah = Ic + Id + Ir - Constant

                    df5 = pd.DataFrame({
                        '1': ['Ic', 'Id', 'Ir', 'Constant', 'Final A/Ah'],
                        '2': [Ic, Id, Ir, Constant, final_a_ah]
                    })

                    # --- Normalized, Standardized, OCV ---
                    target_rows = 55
                    df4_filtered = df3.groupby(tracker)['A/Ah'].mean()
                    dca_combined = pd.concat([df['A/Ah'], df2['A/Ah'], df4_filtered]).tolist()

                    OCV_filtered_81 = raw_df[(raw_df['Step'].astype(str).str.strip() == '81') & (raw_df['Current'] == 0) & (raw_df['Temperature'].notna()) & (raw_df['Cycle'].astype(str).str.strip() == '19')]['Voltage'].dropna()
                    OCV_filtered_30 = raw_df[(raw_df['Step'].astype(str).str.strip() == '30') & (raw_df['Current'] == 0) & (raw_df['Temperature'].notna())]['Voltage'].dropna()
                    OCV_filtered_43 = raw_df[(raw_df['Step'].astype(str).str.strip() == '43') & (raw_df['Current'] == 0) & (raw_df['Temperature'].notna())]['Voltage'].dropna()

                    OCV_Combined = pd.concat([OCV_filtered_30, OCV_filtered_43, OCV_filtered_81]).reset_index(drop=True).iloc[:target_rows].tolist()
                    normalized = [final_a_ah] * target_rows

                    df6 = pd.DataFrame({
                        'Normalized': normalized,
                        'Standard DCA': dca_combined[:target_rows],
                        'OCV': OCV_Combined
                    })

                    # --- Excel Output ---
                    df.to_excel(writer, sheet_name='IC', index=False)
                    df2.to_excel(writer, sheet_name='ID', index=False)
                    df3.to_excel(writer, sheet_name='IR', index=False)
                    IR_Calculation.to_excel(writer, sheet_name='Final', index=False)
                    df4.to_excel(writer, sheet_name='Final', index=False, header=False, startcol=len(IR_Calculation.columns) + 1)
                    df5.to_excel(writer, sheet_name='Final', index=False, header=False, startcol=len(IR_Calculation.columns) + 4)
                    df6.to_excel(writer, sheet_name='Cumulative', index=False)

                    worksheet1 = writer.sheets['IC']
                    worksheet2 = writer.sheets['ID']
                    worksheet3 = writer.sheets['IR']
                    worksheet_final = writer.sheets['Final']
                    worksheet_cum = writer.sheets['Cumulative']

                    # --- Yellow Tables : Summary of each Sheet --- 
                    worksheet_final.conditional_format('A1:B16', {'type': 'no_blanks', 'format': workbook.add_format({'bg_color':'#FFF2CC'})})
                    worksheet_final.conditional_format('D1:E8',  {'type': 'no_blanks', 'format': workbook.add_format({'bg_color': '#E2EFDA'})})
                    worksheet_final.conditional_format('G1:H5',  {'type': 'no_blanks', 'format': workbook.add_format({'bg_color': '#D9E1F2'})})

                    orange_format = workbook.add_format({'bg_color': '#FFC000', 'font_color': '#000000'})
                    orange_num_format = workbook.add_format({'num_format': '0.0000', 'bg_color': '#FFC000', 'font_color': '#000000'})

                    for i, row in df_summary.iterrows():
                        worksheet1.write(1 + i, 21, row['Metric'], orange_format)
                        worksheet1.write(1 + i, 22, row['Value'], orange_num_format)

                    for i, row in df2_summary.iterrows():
                        worksheet2.write(1 + i, 21, row['Metric'], orange_format)
                        worksheet2.write(1 + i, 22, row['Value'], orange_num_format)

                    # --- DCA Charting Layout ---
                    chart1 = workbook.add_chart({'type': 'line'})
                    chart1.add_series({
                        'categories': ['IC', 1, df.columns.get_loc('Cycle'), len(df), df.columns.get_loc('Cycle')],
                        'values':     ['IC', 1, df.columns.get_loc('Current'), len(df), df.columns.get_loc('Current')],
                        'line': {'width': 1.0, 'color': 'blue'},
                        'marker': {'type': 'circle', 'size': 5, 'border': {'color': 'blue'}, 'fill': {'color': 'blue'}}})
                    c1_diff = df['Current'].max() - df['Current'].min()
                    chart1.set_title({'name': f'{battery_code} - Ic (After Discharge History)'})
                    chart1.set_x_axis({'name': '20 Micro cycles', 'label_position': 'low'})
                    chart1.set_y_axis({'name': 'Charge Current (A)', 'min': df['Current'].min() - (c1_diff/5), 'max': df['Current'].max() + (c1_diff/5)})
                    chart1.set_size({'width': 750, 'height': 300})
                    worksheet1.insert_chart('B24', chart1)

                    chart2 = workbook.add_chart({'type': 'line'})
                    chart2.add_series({
                        'categories': ['IC', 1, df.columns.get_loc('Cycle'), len(df), df.columns.get_loc('Cycle')],
                        'values':     ['IC', 1, df.columns.get_loc('A/Ah'), len(df), df.columns.get_loc('A/Ah')],
                        'line': {'width': 1.0, 'color': 'blue'},
                        'marker': {'type': 'circle', 'size': 5, 'border': {'color': 'blue'}, 'fill': {'color': 'blue'}}})
                    c2_diff = df['A/Ah'].max() - df['A/Ah'].min()
                    chart2.set_title({'name': f'{battery_code} - Ic Normalized (After Discharge History)'})
                    chart2.set_x_axis({'name': '20C Micro cycles', 'label_position': 'low'})
                    chart2.set_y_axis({'name': 'Normalized Charge Current (A/Ah)', 'min': df['A/Ah'].min() - (c2_diff/5), 'max': df['A/Ah'].max() + (c2_diff/5)})
                    chart2.set_size({'width': 750, 'height': 300})
                    worksheet1.insert_chart('N24', chart2)

                    chart3 = workbook.add_chart({'type': 'line'})  
                    chart3.add_series({
                        'categories': ['ID', 1, df2.columns.get_loc('Cycle'), len(df2), df2.columns.get_loc('Cycle')],
                        'values':     ['ID', 1, df2.columns.get_loc('Current'), len(df2), df2.columns.get_loc('Current')],
                        'line': {'width': 1.0, 'color': 'blue'},
                        'marker': {'type': 'circle', 'size': 5, 'border': {'color': 'blue'}, 'fill': {'color': 'blue'}}})
                    c3_diff = df2['Current'].max() - df2['Current'].min()
                    chart3.set_title({'name': f'{battery_code} - Id (After Discharge History)'})
                    chart3.set_x_axis({'name': '20 Micro cycles', 'label_position': 'low'})
                    chart3.set_y_axis({'name': 'Charge Current A', 'min': df2['Current'].min() - (c3_diff/5), 'max': df2['Current'].max() + (c3_diff/5)})
                    chart3.set_size({'width': 750, 'height': 300})
                    worksheet2.insert_chart('B24', chart3)

                    chart4 = workbook.add_chart({'type': 'line'}) 
                    chart4.add_series({
                        'categories': ['ID', 1, df2.columns.get_loc('Cycle'), len(df2), df2.columns.get_loc('Cycle')],
                        'values':     ['ID', 1, df2.columns.get_loc('A/Ah'), len(df2), df2.columns.get_loc('A/Ah')],
                        'line': {'width': 1.0, 'color': 'blue'},
                        'marker': {'type': 'circle', 'size': 5, 'border': {'color': 'blue'}, 'fill': {'color': 'blue'}}})
                    c4_diff = df2['A/Ah'].max() - df2['A/Ah'].min()
                    chart4.set_title({'name': f'{battery_code} - Id Normalized (After Discharge History)'})
                    chart4.set_x_axis({'name': '20 Micro cycles', 'label_position': 'low'})
                    chart4.set_y_axis({'name': 'Normalized Charge Current (A/Ah)', 'min': df2['A/Ah'].min() - (c4_diff/5), 'max': df2['A/Ah'].max() + (c4_diff/5)})
                    chart4.set_size({'width': 750, 'height': 300})
                    worksheet2.insert_chart('N24', chart4)

                    chart_ir = workbook.add_chart({'type': 'scatter', 'subtype': 'smooth_with_markers'})
                    cycle_length = 38 
                    total_ir_rows = len(df3)
                    ir_cycle_col = df3.columns.get_loc('Cycle')
                    ir_current_col = df3.columns.get_loc('Current')

                    cycle_num = 1
                    for s_row in range(1, total_ir_rows + 1, cycle_length):
                        e_row = min(s_row + cycle_length - 1, total_ir_rows)
                        chart_ir.add_series({
                            'name':       f'Repetition {cycle_num}',
                            'categories': ['IR', s_row, ir_cycle_col, e_row, ir_cycle_col],
                            'values':     ['IR', s_row, ir_current_col, e_row, ir_current_col],
                            'marker':     {'type': 'circle', 'size': 4},
                            'line':       {'width': 1.0}
                        })
                        cycle_num += 1
                    chart_ir.set_title({'name': f'{battery_code} - IR Real World Start-Stop Regenerative Current'})
                    chart_ir.set_x_axis({'name': 'Cycle', 'min': 0, 'max': 20, 'major_unit': 1})
                    chart_ir.set_y_axis({'name': 'Current (A)'})
                    chart_ir.set_size({'width': 1000, 'height': 500})
                    worksheet3.insert_chart('V2', chart_ir)

                    max_row = len(df6)
                    chart_cum = workbook.add_chart({'type': 'line'})
                    chart_cum.add_series({
                        'name':       ['Cumulative', 0, 0],
                        'values':     ['Cumulative', 1, 0, max_row, 0],
                        'line':       {'color': '#4472C4', 'width': 1.5},
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#4472C4'}, 'fill': {'color': '#4472C4'}}
                    })
                    chart_cum.add_series({
                        'name':       ['Cumulative', 0, 1],
                        'values':     ['Cumulative', 1, 1, max_row, 1],
                        'line':       {'none': True},
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#ED7D31'}, 'fill': {'color': '#ED7D31'}}
                    })
                    chart_cum.add_series({
                        'name':       ['Cumulative', 0, 2],
                        'values':     ['Cumulative', 1, 2, max_row, 2],
                        'line':       {'none': True},
                        'marker':     {'type': 'circle', 'size': 5, 'border': {'color': '#70AD47'}, 'fill': {'color': '#70AD47'}},
                        'y2_axis':    True
                    })
                    chart_cum.set_title({'name': f'{battery_code} - Gen 3 DCA'}) 
                    chart_cum.set_y_axis({'name': 'Normalized Charge current (A/Ah)', 'min': 0.00, 'max': 1.10, 'major_unit': 0.10})
                    chart_cum.set_y2_axis({'name': 'OCV', 'min': 12.50, 'max': 13.30, 'major_unit': 0.10})
                    chart_cum.set_size({'width': 1000, 'height': 500})
                    chart_cum.set_legend({'none': True})
                    worksheet_cum.insert_chart('F4', chart_cum)

                # Test B: DoD 17.5% TEST ENGINE
                elif test_type == "DoD_17.5":
                    battery_name = battery_code
                    df = raw_df.copy()
                    df.rename(columns={df.columns[1]: 'Status'}, inplace=True)

                    # --- Calculate Voltage per Cell and Current per Cell ---
                    df['Voltage'] = pd.to_numeric(df['Voltage'], errors='coerce')
                    df['Current'] = pd.to_numeric(df['Current'], errors='coerce')
                    df['AhStep'] = pd.to_numeric(df['AhStep'], errors='coerce')

                    df['Voltage per Cell'] = df['Voltage'] / 6
                    df['Current per Cell'] = df['Current'] / 6
                    df['SoC'] = 0.00 

                    # --- SoC Computation ---
                    for i in range(len(df)):
                        if i == 0: 
                            first_ah = 0 if pd.isna(df.loc[i, 'AhStep']) else df.loc[i, 'AhStep']
                            if df.loc[i, 'Status'] == 'CHA':
                                df.loc[i, 'SoC'] = (60 + first_ah) / 60
                            else:
                                df.loc[i, 'SoC'] = (60 - first_ah) / 60
                        else: 
                            previous_SoC = df.loc[i-1, 'SoC']
                            ahstep = df.loc[i, 'AhStep']
                            ahstep_before = df.loc[i-1, 'AhStep']
                            diff_ahstep = ahstep - ahstep_before
                            
                            if pd.isna(ahstep) or pd.isna(ahstep_before) or ahstep == ahstep_before:
                                diff_ahstep = 0
                                
                            if df.loc[i, 'Status'] == 'CHA':
                                df.loc[i, 'SoC'] = previous_SoC + (diff_ahstep / 60)    
                            else:
                                df.loc[i, 'SoC'] = previous_SoC - (diff_ahstep / 60)
                                
                    df['SoC%'] = df['SoC'] / 1.0 * 100 
                    start_discharge = (df['SoC%'] < 100).idxmax() 
                    full_charge = (df.loc[start_discharge:, 'SoC%'] >= 100).idxmax() 
                    
                    if full_charge > start_discharge:
                        df = df.loc[:full_charge]

                    # --- Filtering Sheet 2 ---
                    df2 = raw_df.copy()
                    df2.rename(columns={df2.columns[1]: 'Status'}, inplace=True) 
                    df2 = df2[df2['Status'] == 'DCH']
                    df2 = df2[df2['Step'].isin([3, 7])]
                    df2 = df2[(df2['Step Time'] = '00:30:00.000') & (df2['Step Time'] < '00:30:00.1')]
                    
                    # --- Different Names for Temperatures --- 
                    temp_col = 'Temperature_1' if 'Temperature_1' in df2.columns else 'Temperature'
                    df2 = df2.dropna(subset=[temp_col])
                    
                    df2 = df2[df2['Cycle'] != df2['Cycle'].shift()]
                    df2['Total Cycle'] = range(0, len(df2)) 

                    # --- Excel Output ---
                    df.to_excel(writer, sheet_name='Sheet 1', index=False)
                    df2.to_excel(writer, sheet_name='Sheet 2', index=False)

                    worksheet1 = writer.sheets['Sheet 1']
                    worksheet2 = writer.sheets['Sheet 2']   

                    # --- DoD17.5% Charting Layout ---
                    chart1 = workbook.add_chart({'type': 'line'})
                    chart1.add_series({
                        'categories': ['Sheet 1', 1, df.columns.get_loc('Prog Time'), len(df), df.columns.get_loc('Prog Time')],
                        'values':     ['Sheet 1', 1, df.columns.get_loc('Voltage per Cell'), len(df), df.columns.get_loc('Voltage per Cell')],
                        'line': {'width': 0.01, 'color': 'blue'}})
                    chart1.set_title({'name': f'{battery_name} - Voltage vs Time'})
                    chart1.set_x_axis({'name': 'Prog Time', 'label_position': 'low'})
                    chart1.set_y_axis({'name': 'Voltage per Cell', 'min': 1.5, 'max': 3})
                    chart1.set_size({'width': 1000, 'height': 500})
                    worksheet1.insert_chart('V2', chart1)

                    chart2 = workbook.add_chart({'type': 'line'})
                    chart2.add_series({
                        'categories': ['Sheet 1', 1, df.columns.get_loc('Prog Time'), len(df), df.columns.get_loc('Prog Time')],
                        'values':     ['Sheet 1', 1, df.columns.get_loc('Current per Cell'), len(df), df.columns.get_loc('Current per Cell')],
                        'line': {'width': 0.01, 'color': 'blue'}})
                    chart2.set_title({'name': f'{battery_name} - Current vs Time'})
                    chart2.set_x_axis({'name': 'Prog Time', 'label_position': 'low'})
                    chart2.set_y_axis({'name': 'Current per Cell', 'min': -4, 'max': 4})
                    chart2.set_size({'width': 1000, 'height': 500})
                    worksheet1.insert_chart('V28', chart2)

                    chart3 = workbook.add_chart({'type': 'line'})  
                    chart3.add_series({
                        'categories': ['Sheet 1', 1, df.columns.get_loc('Prog Time'), len(df), df.columns.get_loc('Prog Time')],
                        'values':     ['Sheet 1', 1, df.columns.get_loc('SoC%'), len(df), df.columns.get_loc('SoC%')],
                        'line': {'width': 0.01, 'color': 'blue'}})
                    chart3.set_title({'name': f'{battery_name} - State of Charge (SoC%) vs Time'})
                    chart3.set_x_axis({'name': 'Prog Time', 'label_position': 'low'})
                    chart3.set_y_axis({'name': 'SoC%', 'min': 0, 'max': 100})
                    chart3.set_size({'width': 1000, 'height': 500})
                    worksheet1.insert_chart('V54', chart3)

                    chart4 = workbook.add_chart({'type': 'scatter', 'subtype': 'straight_with_markers'})
                    chart4.add_series({
                        'categories': ['Sheet 2', 1, df2.columns.get_loc('Total Cycle'), len(df2), df2.columns.get_loc('Total Cycle')],
                        'values':     ['Sheet 2', 1, df2.columns.get_loc('Voltage'), len(df2), df2.columns.get_loc('Voltage')],
                        'line': {'width': 0.01, 'color': 'blue'}})
                    chart4.set_title({'name': f'17.5% DoD of {battery_name}'})  
                    chart4.set_x_axis({'name': 'Cycle', 'min': 0, 'label_position': 'low', 'major_unit': 85})
                    chart4.set_y_axis({'name': 'Voltage', 'min': 10, 'max': 12.5})
                    chart4.set_size({'width': 1000, 'height': 500})
                    worksheet2.insert_chart('U2', chart4)

                # Test C: Other types
                else:
                    st.warning(f"Bypassing custom calculations: Visual sheets layout for '{test_type}' hasn't been defined yet.")
                    st.stop()

            # --- Streamlit Download Button ---
            st.success(f"✅ Data processed successfully for {test_type} test!")
            
            st.download_button(
                label=f"📥 Download {output_filename}",
                data=output_buffer.getvalue(),
                file_name=output_filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"An error occurred while processing the file: {e}")
